import sys
from pathlib import Path
import openpyxl

# Ensure UTF-8 output encoding for emojis on Windows console
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Add backend directory to sys.path if not present
BACKEND_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BACKEND_DIR.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from database import get_db_connection, init_db

def seed_events():
    """Seeds events table from data/events_seed.xlsx if the table is currently empty."""
    init_db()
    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if table already has rows
    cursor.execute("SELECT COUNT(*) AS cnt FROM events;")
    count = cursor.fetchone()["cnt"]
    if count > 0:
        print(f"Events table already has {count} rows. Skipping seeding.")
        conn.close()
        return count

    # Check if data/events_seed.xlsx exists
    xlsx_path = PROJECT_ROOT / "data" / "events_seed.xlsx"
    if not xlsx_path.exists():
        conn.close()
        raise FileNotFoundError(
            f"Required seed file not found at: {xlsx_path}. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    # Load workbook and read 'Events' sheet
    wb = openpyxl.load_workbook(filename=str(xlsx_path))
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError("Spreadsheet does not contain a sheet named 'Events'.")

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        conn.close()
        raise ValueError("Events sheet is empty or contains no data rows.")

    # Row 1 is header row: id, title, category, description, event_date, event_time, venue, capacity, fee, icon
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    expected_headers = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):  # Skip completely empty trailing rows
            continue
        ev_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            int(ev_id),
            str(title).strip(),
            str(category).strip(),
            str(description).strip() if description else "",
            str(event_date).strip(),
            str(event_time).strip(),
            str(venue).strip(),
            int(capacity),
            int(fee) if fee is not None else 0,
            str(icon).strip() if icon else "🎉"
        ))

    conn.commit()
    cursor.execute("SELECT COUNT(*) AS cnt FROM events;")
    final_count = cursor.fetchone()["cnt"]
    conn.close()
    print(f"Successfully seeded {final_count} events into the database.")
    return final_count

if __name__ == "__main__":
    count = seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events;")
    all_events = cursor.fetchall()
    print(f"\n--- Events in Database ({len(all_events)}) ---")
    for ev in all_events:
        print(f"[{ev['id']}] {ev['icon']} {ev['title']} ({ev['category']}) | {ev['event_date']} {ev['event_time']} | Venue: {ev['venue']} | Cap: {ev['capacity']} | Fee: ₹{ev['fee']}")
    conn.close()
